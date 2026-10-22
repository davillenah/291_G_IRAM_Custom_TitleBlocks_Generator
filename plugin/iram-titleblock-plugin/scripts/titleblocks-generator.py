import subprocess, os, html, sys
from openpyxl import load_workbook
W = os.getcwd(); HERE = os.path.dirname(os.path.abspath(__file__)); GEN = os.path.join(HERE, "sheet-generator.py")
# Dos formas de uso:
#  - carpeta del proyecto con input-titleblocks.xlsx (plugin): logos y salida en esa misma carpeta (svg/)
#  - estructura del proyecto: src/ (Excel y scripts), assets/ (logos), test/ (salida)
if os.path.exists(os.path.join(W, "input-titleblocks.xlsx")):
    INPUT, LOGOS, OUT = W, W, os.path.join(W, "svg")
else:
    ROOT = os.path.dirname(HERE)
    INPUT, LOGOS, OUT = HERE, os.path.join(ROOT, "assets"), os.path.join(ROOT, "test")
ENV = dict(os.environ, IRAM_PROYECTO=LOGOS)  # el generador busca los logos en la carpeta del proyecto
KEYS = ["id","formato","proy_fecha","proy_nombre","ver_fecha","ver_nombre","val_fecha","val_nombre",
        "empresa","empresa_cuit","empresa_email","empresa_web","cliente","cliente_cuit","cliente_email","cliente_web",
        "proyecto","descripcion","ubicacion","tipo_plano","denominacion","denominacion_2","escala",
        "reemplaza","reemplazado_por","plano_n","pagina","estado","revision","mod","propio","archivo"]
wb = load_workbook(os.path.join(INPUT, "input-titleblocks.xlsx"), data_only=True); ws = wb["Láminas"]
os.makedirs(OUT, exist_ok=True)
for f in os.listdir(OUT):
    if f.endswith((".svg", ".dxf")): os.remove(os.path.join(OUT, f))
IDX = {k: i + 1 for i, k in enumerate(KEYS)}
for r in range(2, ws.max_row + 1):
    row = {k: ws.cell(r, IDX[k]).value for k in KEYS}
    if not row["formato"]: continue
    mod = row["mod"] == "SI"
    dst = os.path.join(OUT, row["archivo"])
    subprocess.run(["python3", "-I", GEN, row["formato"], dst] + (["--mod"] if mod else []),
                   check=True, capture_output=True, cwd=OUT, env=ENV)
    s = open(dst, encoding="utf-8").read()
    for k in KEYS:
        if k in ("id", "mod", "propio", "archivo"): continue
        v = row[k]
        v = "" if v is None else (v.strftime("%d/%m/%Y") if hasattr(v, "strftime") else str(v))
        s = s.replace("{%s}" % k, html.escape(v, quote=False))
    s = s.replace("{plano_propio}", html.escape(row["propio"] or "", quote=False))
    open(dst, "w", encoding="utf-8").write(s)
    print("regen", row["archivo"], "mod" if mod else "")

# DXF de cada lámina, para CAD (mismas medidas en mm)
sys.path.insert(0, HERE)
from svg_to_dxf import convertir
for f in sorted(os.listdir(OUT)):
    if f.startswith("sheet-") and f.endswith(".svg"):
        convertir(os.path.join(OUT, f))
